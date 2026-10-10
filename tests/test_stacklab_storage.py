"""Portable storage, report, and independent example reference checks."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from stacklab.compiler import ModelError, compile_project
from stacklab.examples import create_example, write_example_projects
from stacklab.persistence import (
    ProjectFormatError, load_project, load_project_bundle, save_project,
)
from stacklab.reporting import export_csv, export_pdf, export_xlsx, report_rows
from stacklab.solvers import analyze


@pytest.mark.parametrize("example", list("ABCDE"))
def test_example_project_round_trip(tmp_path: Path, example: str) -> None:
    original = create_example(example)
    destination = tmp_path / f"{example}.stack1d"
    presentation = {"zoom": 1.25, "selected": "drawing"}
    save_project(destination, original, presentation)
    bundle = load_project_bundle(destination)
    assert bundle.project.to_dict() == original.to_dict()
    assert bundle.presentation == presentation
    assert load_project(destination).id == original.id
    with ZipFile(destination) as archive:
        assert set(archive.namelist()) == {
            "manifest.json", "engineering.json", "presentation.json",
        }
        engineering = json.loads(archive.read("engineering.json"))
        assert "view" not in engineering
        assert engineering["name"] == original.name


def test_save_failure_preserves_previous_archive(tmp_path: Path) -> None:
    destination = tmp_path / "assembly.stack1d"
    save_project(destination, create_example("A"))
    before = destination.read_bytes()
    invalid = create_example("A")
    invalid.instances[0].id = invalid.definitions[0].id  # duplicate stable ID
    with pytest.raises((ValueError, ProjectFormatError)):
        save_project(destination, invalid)
    assert destination.read_bytes() == before


def test_malformed_archive_and_schema_are_rejected(tmp_path: Path) -> None:
    destination = tmp_path / "bad.stack1d"
    destination.write_text("not a ZIP archive", encoding="utf-8")
    with pytest.raises(ProjectFormatError, match="Cannot open"):
        load_project(destination)
    with ZipFile(destination, "w") as archive:
        archive.writestr("manifest.json", json.dumps({"format": "stacklab-1d", "schema_version": 999}))
    with pytest.raises(ProjectFormatError, match="Unsupported"):
        load_project(destination)
    with ZipFile(destination, "w") as archive:
        archive.writestr("manifest.json", json.dumps({"format": "stacklab-1d", "schema_version": 2}))
        archive.writestr("engineering.json", "{}")
        archive.writestr("presentation.json", "{}")
    with pytest.raises(ProjectFormatError, match="requires a non-empty id"):
        load_project(destination)


def test_version_one_project_migrates(tmp_path: Path) -> None:
    destination = tmp_path / "older.stack1d"
    project = create_example("A")
    with ZipFile(destination, "w") as archive:
        archive.writestr("manifest.json", json.dumps({"format": "stacklab-1d", "schema_version": 1}))
        archive.writestr("project.json", json.dumps({"engineering": project.to_dict(),
                                                     "presentation": {"zoom": 2.0}}))
    bundle = load_project_bundle(destination)
    assert bundle.project.to_dict() == project.to_dict()
    assert bundle.presentation == {"zoom": 2.0}


def test_reference_fixed_chain_and_floating_assemblies() -> None:
    fixed = analyze(create_example("A"), "a-gap", methods=("worst_case",))
    assert fixed.chain.equation
    assert fixed.worst_case.nominal == pytest.approx(1.0)
    assert fixed.worst_case.minimum == pytest.approx(0.55)
    assert fixed.worst_case.maximum == pytest.approx(1.45)

    floating = create_example("B")
    seated = analyze(floating, "b-clearance", methods=("worst_case",)).worst_case
    assert seated.nominal == pytest.approx(35.0)
    assert seated.minimum == pytest.approx(34.55)
    assert seated.maximum == pytest.approx(35.45)
    free = analyze(floating, "b-free-gap", methods=("worst_case",)).worst_case
    assert free.nominal is None
    assert free.minimum == pytest.approx(0.0)
    assert free.maximum == pytest.approx(35.45)
    right = analyze(floating, "b-right-gap", methods=("worst_case",)).worst_case
    assert (right.nominal, right.minimum, right.maximum) == pytest.approx((0.0, 0.0, 0.0))
    centered = analyze(floating, "b-centered-gap", methods=("worst_case",)).worst_case
    assert centered.nominal == pytest.approx(17.5)
    assert centered.minimum == pytest.approx(17.275)
    assert centered.maximum == pytest.approx(17.725)


def test_coupled_floating_parts_and_conflicting_loop() -> None:
    coupled = create_example("C")
    free = analyze(coupled, "c-interblock", methods=("worst_case",)).worst_case
    assert free.nominal is None
    assert free.minimum == pytest.approx(0.0)
    assert free.maximum == pytest.approx(50.45)
    seated = analyze(coupled, "c-rear", methods=("worst_case",)).worst_case
    assert seated.nominal == pytest.approx(50.0)
    assert seated.minimum == pytest.approx(49.55)
    assert seated.maximum == pytest.approx(50.45)
    with pytest.raises(ModelError, match="feasible|inconsistent|constraint"):
        compile_project(create_example("D"))


def test_statistical_example_matches_independent_variance() -> None:
    project = create_example("E")
    result = analyze(project, "e-gap", methods=("rss", "monte_carlo"),
                     samples=30_000, seed=2026)
    expected_variance = (0.20 / 3) ** 2 + (0.10 / 3) ** 2 + (0.15 / 3) ** 2
    assert result.rss.mean == pytest.approx(1.0)
    assert result.rss.variance == pytest.approx(expected_variance)
    assert result.monte_carlo.valid_samples == 30_000
    # Mean SE < 0.001; standard-deviation SE < 0.0004 at this sample count.
    assert result.monte_carlo.mean == pytest.approx(1.0, abs=0.004)
    assert result.monte_carlo.std == pytest.approx(expected_variance ** 0.5, abs=0.002)


def test_reports_use_computed_result_and_all_formats(tmp_path: Path) -> None:
    project = create_example("A")
    result = analyze(project, "a-gap", methods=("worst_case",))
    rows = report_rows(project, result)
    assert any(row.item == "Minimum gap" and row.value == "0.55" for row in rows)
    assert any(row.item == "Maximum gap" and row.value == "1.45" for row in rows)

    csv_path, xlsx_path, pdf_path = (tmp_path / f"report.{ext}"
                                     for ext in ("csv", "xlsx", "pdf"))
    export_csv(csv_path, project, result)
    export_xlsx(xlsx_path, project, result)
    export_pdf(pdf_path, project, result)
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    assert any(row["Item"] == "Minimum gap" and row["Value"] == "0.55" for row in csv_rows)
    workbook = load_workbook(xlsx_path, read_only=True)
    assert any(row[2] == "Maximum gap" and row[3] == "1.45"
               for row in workbook.active.iter_rows(min_row=2, values_only=True))
    workbook.close()
    assert pdf_path.read_bytes().startswith(b"%PDF")


def test_spreadsheet_exports_keep_user_labels_as_text(tmp_path: Path) -> None:
    project = create_example("A")
    project.requirements[0].name = '=HYPERLINK("https://example.invalid")'
    result = analyze(project, "a-gap", methods=("worst_case",))
    csv_path = tmp_path / "safe.csv"
    xlsx_path = tmp_path / "safe.xlsx"
    export_csv(csv_path, project, result)
    export_xlsx(xlsx_path, project, result)
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        assert next(csv.DictReader(handle))["Requirement"].startswith("'=")
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook.active["A2"].data_type == "s"
    workbook.close()


def test_checked_in_examples_are_current(tmp_path: Path) -> None:
    generated = write_example_projects(tmp_path)
    assert len(generated) == 5
    for path in generated:
        checked_in = Path(__file__).resolve().parents[1] / "examples" / path.name
        assert load_project(checked_in).to_dict() == load_project(path).to_dict()
