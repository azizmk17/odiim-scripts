"""Engineering reports generated from a project and actual analysis results."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Mapping
from xml.sax.saxutils import escape

from .domain import AnalysisResult, FaceRef, Project


@dataclass(frozen=True)
class ReportRow:
    requirement: str
    section: str
    item: str
    value: str
    unit: str = ""


def _number(value: float | int | None) -> str:
    if value is None:
        return "Unavailable"
    if isinstance(value, int):
        return str(value)
    return f"{value:.6g}"


def _face_name(project: Project, ref: FaceRef) -> str:
    instance = next((item for item in project.instances if item.id == ref.instance_id), None)
    if instance is None:
        return f"{ref.instance_id}:{ref.face_id}"
    definition = next((item for item in project.definitions if item.id == instance.definition_id), None)
    face = next((item for item in definition.faces if item.id == ref.face_id), None) if definition else None
    return f"{instance.name}.{face.name if face else ref.face_id}"


def _results_list(results: AnalysisResult | Iterable[AnalysisResult] |
                  Mapping[str, AnalysisResult]) -> list[AnalysisResult]:
    if isinstance(results, AnalysisResult):
        normalized = [results]
    elif isinstance(results, Mapping):
        normalized = list(results.values())
    else:
        normalized = list(results)
    if not normalized or not all(isinstance(result, AnalysisResult) for result in normalized):
        raise ValueError("Reports require one or more completed AnalysisResult objects")
    return normalized


def report_rows(project: Project, results: AnalysisResult | Iterable[AnalysisResult] |
                Mapping[str, AnalysisResult]) -> list[ReportRow]:
    """Flatten real analysis data into rows shared by PDF, CSV, and XLSX."""
    output: list[ReportRow] = []
    sources = {source.id: source for source in project.sources}
    dimensions = {dimension.id: dimension for dimension in project.dimensions}
    requirements = {requirement.id: requirement for requirement in project.requirements}
    policies = {policy.id: policy for policy in project.policies}

    for result in _results_list(results):
        requirement = requirements.get(result.requirement_id)
        if requirement is None:
            raise ValueError(f"Analysis result refers to missing requirement {result.requirement_id}")
        name = requirement.name

        def add(section: str, item: str, value: object, unit: str = "") -> None:
            if isinstance(value, (float, int)) and not isinstance(value, bool):
                formatted = _number(value)
            elif value is None:
                formatted = "Unavailable"
            elif isinstance(value, bool):
                formatted = "Yes" if value else "No"
            else:
                formatted = str(value)
            output.append(ReportRow(name, section, item, formatted, unit))

        add("Requirement", "First face", _face_name(project, requirement.first))
        add("Requirement", "Second face", _face_name(project, requirement.second))
        add("Requirement", "Measurement direction", requirement.direction)
        add("Requirement", "Minimum acceptable", requirement.min_value, project.unit)
        add("Requirement", "Maximum acceptable", requirement.max_value, project.unit)
        policy = policies.get(requirement.policy_id) if requirement.policy_id else None
        add("Assembly", "Positioning policy", f"{policy.name} ({policy.kind})" if policy else "Free")
        for constraint in project.constraints:
            relation = constraint.kind
            if constraint.first and constraint.second:
                relation += f" {_face_name(project, constraint.first)} → {_face_name(project, constraint.second)}"
            elif constraint.first:
                relation += f" {_face_name(project, constraint.first)}"
            elif constraint.instance_id:
                relation += f" {constraint.instance_id}"
            if constraint.kind in {"fixed_position", "fixed_face", "fixed_offset"}:
                relation += f" = {_number(constraint.value)} {project.unit}"
            elif constraint.lower is not None or constraint.upper is not None:
                relation += f" [{_number(constraint.lower)}, {_number(constraint.upper)}] {project.unit}"
            add("Assembly", constraint.name, relation)
        for contact in project.contacts:
            add("Assembly", contact.name or contact.id,
                f"Nonpenetration: {_face_name(project, contact.second)} ≥ {_face_name(project, contact.first)}")

        for dimension in project.dimensions:
            detail = f"{_face_name(project, dimension.first)} → {_face_name(project, dimension.second)}"
            tolerance = f"{_number(dimension.nominal)} +{_number(dimension.tolerance.upper)} / {_number(dimension.tolerance.lower)}"
            source = sources.get(dimension.source_id) if dimension.source_id else None
            if source:
                tolerance += f"; source {source.name}"
            add("Inputs", f"{dimension.name} ({dimension.kind}; {detail})", tolerance, project.unit)
        for source in project.sources:
            distribution = f"{source.distribution}; limits [{_number(source.lower)}, {_number(source.upper)}]"
            if source.std is not None:
                distribution += f"; σ={_number(source.std)}"
            add("Inputs", f"Variation source {source.name}", distribution, project.unit)
        for correlation in project.correlations:
            add("Inputs", "Correlation",
                f"{correlation.first_source_id} ↔ {correlation.second_source_id}: ρ={_number(correlation.rho)}")

        chain = result.chain
        add("Dimensional chain", "Equation", chain.equation or "Unavailable")
        add("Dimensional chain", "Floating assembly", chain.floating)
        add("Dimensional chain", "Translation degrees of freedom", chain.degrees_of_freedom)
        def worst_case_impact(term) -> float:
            source = sources.get(term.source_id)
            if source is None:
                return abs(term.coefficient)
            return abs(term.coefficient) * max(abs(source.lower), abs(source.upper))

        for term in sorted(chain.terms, key=worst_case_impact, reverse=True):
            source = sources.get(term.source_id)
            label = source.name if source else term.source_id
            add("Contributors", f"{label}: sensitivity", term.coefficient)
            if source:
                lower_effect = min(term.coefficient * source.lower, term.coefficient * source.upper)
                upper_effect = max(term.coefficient * source.lower, term.coefficient * source.upper)
                add("Contributors", f"{label}: tolerance effect",
                    f"[{_number(lower_effect)}, {_number(upper_effect)}]", project.unit)
            if term.dimension_ids:
                names = [dimensions[identifier].name if identifier in dimensions else identifier
                         for identifier in term.dimension_ids]
                add("Contributors", f"{label}: dimensions", ", ".join(names))

        worst_case = result.worst_case
        if worst_case is not None:
            add("Worst case — exact feasible envelope", "Nominal gap", worst_case.nominal, project.unit)
            add("Worst case — exact feasible envelope", "Minimum gap", worst_case.minimum, project.unit)
            add("Worst case — exact feasible envelope", "Maximum gap", worst_case.maximum, project.unit)
            add("Worst case — exact feasible envelope", "Interference possible", worst_case.possible_interference)
            add("Worst case — exact feasible envelope", "Meets specified limits", worst_case.meets_limits)
            add("Worst case — exact feasible envelope", "Policy", worst_case.policy_kind)
            for label, state in (("Minimum", worst_case.minimum_state), ("Maximum", worst_case.maximum_state)):
                add("Extreme assembly states", f"{label} active contacts",
                    ", ".join(next((contact.name for contact in project.contacts
                                    if contact.id == contact_id and contact.name), contact_id)
                              for contact_id in state.active_contacts) or "None")
                for instance_id, position in state.translations.items():
                    add("Extreme assembly states", f"{label} translation {instance_id}", position, project.unit)
                for source_id, deviation in state.source_values.items():
                    add("Extreme manufacturing states", f"{label} deviation {source_id}", deviation, project.unit)
            for warning in worst_case.warnings:
                add("Warnings and limits", "Worst case", warning)
        rss = result.rss
        if rss is not None:
            add("RSS — affine statistical model", "Mean gap", rss.mean, project.unit)
            add("RSS — affine statistical model", "Variance", rss.variance, f"{project.unit}²")
            add("RSS — affine statistical model", "Standard deviation", rss.std, project.unit)
            add("RSS — affine statistical model", "Sigma level", rss.sigma_level)
            add("RSS — affine statistical model", "Lower report limit", rss.lower, project.unit)
            add("RSS — affine statistical model", "Upper report limit", rss.upper, project.unit)
            for source_id, contribution in rss.variance_contributions.items():
                add("RSS — affine statistical model", f"Variance allocation {source_id}",
                    contribution, f"{project.unit}²")
            for warning in rss.warnings:
                add("Warnings and limits", "RSS", warning)
        monte_carlo = result.monte_carlo
        if monte_carlo is not None:
            add("Monte Carlo — statistical estimate", "Requested samples", monte_carlo.requested_samples)
            add("Monte Carlo — statistical estimate", "Valid samples", monte_carlo.valid_samples)
            add("Monte Carlo — statistical estimate", "Infeasible assemblies", monte_carlo.infeasible_samples)
            add("Monte Carlo — statistical estimate", "Random seed", monte_carlo.seed)
            for field_name in ("minimum", "maximum", "mean", "std"):
                add("Monte Carlo — statistical estimate", field_name.title(),
                    getattr(monte_carlo, field_name), project.unit)
            for percentile, value in monte_carlo.percentiles.items():
                add("Monte Carlo — statistical estimate", f"Percentile {percentile}", value, project.unit)
            add("Monte Carlo — statistical estimate", "Out-of-specification samples",
                monte_carlo.specification_failures)
            add("Monte Carlo — statistical estimate", "Conditional failure probability",
                monte_carlo.conditional_failure_probability)
            add("Monte Carlo — statistical estimate", "Overall failure probability",
                monte_carlo.overall_failure_probability)
            add("Monte Carlo — statistical estimate", "PPM outside specification",
                monte_carlo.ppm_outside_spec)
            if monte_carlo.failure_probability_ci95:
                low, high = monte_carlo.failure_probability_ci95
                add("Monte Carlo — statistical estimate", "95% failure probability interval",
                    f"[{_number(low)}, {_number(high)}]")
            for warning in monte_carlo.warnings:
                add("Warnings and limits", "Monte Carlo", warning)
        for diagnostic in result.diagnostics:
            add("Warnings and limits", diagnostic.code,
                f"{diagnostic.severity}: {diagnostic.message}")
    return output


def export_csv(path: str | Path, project: Project,
               results: AnalysisResult | Iterable[AnalysisResult] | Mapping[str, AnalysisResult]) -> None:
    def safe_cell(value: str) -> str:
        # Spreadsheet applications interpret these prefixes as formulas, even
        # when the value came from a user-entered feature name or warning.
        if value and value[0] in "=+-@":
            try:
                float(value)
            except ValueError:
                return "'" + value
        return value

    rows = report_rows(project, results)
    with Path(path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("Requirement", "Section", "Item", "Value", "Unit"))
        for row in rows:
            writer.writerow(tuple(safe_cell(value) for value in
                                  (row.requirement, row.section, row.item, row.value, row.unit)))


def export_xlsx(path: str | Path, project: Project,
                results: AnalysisResult | Iterable[AnalysisResult] | Mapping[str, AnalysisResult]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rows = report_rows(project, results)

    def safe_cell(value: str) -> str:
        return "'" + value if value and value[0] == "=" else value

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Analysis"
    sheet.append(("Requirement", "Section", "Item", "Value", "Unit"))
    for row in rows:
        sheet.append(tuple(safe_cell(value) for value in
                           (row.requirement, row.section, row.item, row.value, row.unit)))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="203656")
    for column, width in {"A": 26, "B": 32, "C": 42, "D": 60, "E": 12}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def export_pdf(path: str | Path, project: Project,
               results: AnalysisResult | Iterable[AnalysisResult] | Mapping[str, AnalysisResult]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .compiler import ModelError
    from .solvers import UnsupportedAnalysis, solve_nominal_geometry

    rows = report_rows(project, results)
    try:
        nominal_positions = solve_nominal_geometry(project)
        sketch_error = None
    except (ModelError, UnsupportedAnalysis) as exc:
        nominal_positions = {}
        sketch_error = str(exc)
    styles = getSampleStyleSheet()
    styles["Normal"].fontSize = 8
    styles["Normal"].leading = 11

    class AssemblySketch(Flowable):
        def __init__(self) -> None:
            super().__init__()
            self.width = 170 * mm
            self.height = max(32 * mm, (len(project.instances) * 25 + 24) * mm)

        def draw(self) -> None:
            plotted: list[tuple[str, list[tuple[str, float]]]] = []
            positions: list[float] = []
            for instance in project.instances:
                definition = next((item for item in project.definitions
                                   if item.id == instance.definition_id), None)
                if definition is None or not instance.visible:
                    continue
                faces = [(face.name, nominal_positions[(instance.id, face.id)])
                         for face in definition.faces
                         if (instance.id, face.id) in nominal_positions]
                if faces:
                    plotted.append((instance.name, faces))
                    positions.extend(value for _, value in faces)
            if not positions:
                message = ("Nominal sketch unavailable: " + sketch_error) if sketch_error else "No sketch geometry in project"
                self.canv.drawString(5, self.height - 15, message[:100])
                return
            lo, hi = min(positions), max(positions)
            span = max(hi - lo, 1.0)
            left, right = 92, self.width - 12
            scale = (right - left) / span
            self.canv.setFont("Helvetica", 7)
            self.canv.setStrokeColor(colors.HexColor("#52677E"))
            for index, (name, faces) in enumerate(plotted):
                y = self.height - (index + 1) * 25 * mm
                self.canv.drawString(0, y, name[:24])
                xs = [left + (value - lo) * scale for _, value in faces]
                self.canv.setFillColor(colors.HexColor("#DFEAF3"))
                self.canv.rect(min(xs), y - 3 * mm, max(max(xs) - min(xs), 1.2 * mm),
                               6 * mm, stroke=1, fill=1)
                self.canv.setFillColor(colors.HexColor("#172C44"))
                occupied: list[list[tuple[float, float]]] = [[], [], []]
                for face_name, value in sorted(faces, key=lambda item: item[1]):
                    x = left + (value - lo) * scale
                    self.canv.line(x, y - 4 * mm, x, y + 4 * mm)
                    label = face_name[:18]
                    label_width = self.canv.stringWidth(label, "Helvetica", 7)
                    label_left = x + 3
                    if label_left + label_width > self.width - 2:
                        label_left = x - label_width - 3
                    interval = (label_left - 2, label_left + label_width + 2)
                    tier = next((tier_index for tier_index, intervals in enumerate(occupied)
                                 if all(interval[1] < start or interval[0] > end
                                        for start, end in intervals)), 2)
                    occupied[tier].append(interval)
                    label_y = y + (7 + 6 * tier) * mm
                    self.canv.setStrokeColor(colors.HexColor("#8899A9"))
                    self.canv.line(x, y + 4 * mm, x, label_y - 1.5 * mm)
                    self.canv.setStrokeColor(colors.HexColor("#52677E"))
                    self.canv.drawString(label_left, label_y, label)
            self.canv.drawRightString(self.width, 4, f"Nominal model sketch ({project.unit})")

    document = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm,
                                 leftMargin=18 * mm, topMargin=17 * mm, bottomMargin=17 * mm)
    story: list = [Paragraph("StackLab 1D engineering analysis", styles["Title"]),
                   Spacer(1, 4 * mm),
                   Paragraph(f"Project: {escape(project.name)} · ID: {escape(project.id)}",
                             styles["Normal"]),
                   Paragraph(f"Analysis exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
                             f" · Units: {escape(project.unit)}", styles["Normal"]),
                   Spacer(1, 5 * mm), Paragraph("Assembly sketch", styles["Heading2"]),
                   AssemblySketch(), Spacer(1, 5 * mm)]
    last_key: tuple[str, str] | None = None
    section_rows: list[ReportRow] = []

    def flush() -> None:
        if not section_rows:
            return
        heading = f"{section_rows[0].requirement} — {section_rows[0].section}"
        story.append(Paragraph(escape(heading), styles["Heading3"]))
        table_data = [["Item", "Value", "Unit"]]
        for row in section_rows:
            table_data.append([Paragraph(escape(row.item), styles["Normal"]),
                               Paragraph(escape(row.value), styles["Normal"]),
                               Paragraph(escape(row.unit), styles["Normal"])])
        table = Table(table_data, colWidths=[58 * mm, 94 * mm, 17 * mm], repeatRows=1,
                      hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#203656")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D9E0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6F9")]),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend((table, Spacer(1, 3 * mm)))

    for row in rows:
        key = row.requirement, row.section
        if last_key is not None and key != last_key:
            flush()
            section_rows = []
        section_rows.append(row)
        last_key = key
    flush()
    document.build(story)
